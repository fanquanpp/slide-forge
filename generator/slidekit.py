# -*- coding: utf-8 -*-
"""SlideForge 渲染引擎：形状/文本工具 + 24 种可复用版式原型 + 切换动效。

所有页面只含版式骨架与占位示例文本，可自由编辑，不含正式内容。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

SW, SH = 13.333, 7.5
import anim as _anim
_VALID_TRANS = set()


def C(h):
    return RGBColor.from_string(h)


def _set_ea(run, ea):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", ea)


def set_alpha(shape, alpha):
    try:
        spPr = shape._element.spPr
        solid = spPr.find(qn("a:solidFill"))
        if solid is None:
            return
        for srgb in solid.iter(qn("a:srgbClr")):
            a = srgb.makeelement(qn("a:alpha"), {"val": str(int(round(alpha * 100000)))})
            srgb.append(a)
    except Exception:
        pass


def set_transition(slide, spec):
    """给单页注入切换特效（由 anim 引擎提供多样化类型）。"""
    _anim.set_transition(slide, spec)


class Deck:
    def __init__(self, theme):
        self.t = theme
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self._page = 0
        self._trans = theme.get("transition", "fade")

    # ---------- 基础 ----------
    def slide(self):
        return self.prs.slides.add_slide(self.blank)

    def bg(self, s, color=None):
        self.rect(s, 0, 0, SW, SH, color or self.t["bg"])

    def rect(self, s, l, tp, w, h, fill=None, line=None, lw=0.75, radius=0.0, alpha=None, shape=MSO_SHAPE.RECTANGLE):
        shp = s.shapes.add_shape(shape, Inches(l), Inches(tp), Inches(w), Inches(h))
        if fill:
            shp.fill.solid(); shp.fill.fore_color.rgb = C(fill)
            if alpha is not None:
                set_alpha(shp, alpha)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = C(line); shp.line.width = Pt(lw)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if radius and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                shp.adjustments[0] = radius
            except Exception:
                pass
        return shp

    def rrect(self, s, l, tp, w, h, fill=None, line=None, lw=0.75, radius=None, alpha=None):
        r = self.t.get("radius", 0.1) if radius is None else radius
        return self.rect(s, l, tp, w, h, fill, line, lw, r, alpha, MSO_SHAPE.ROUNDED_RECTANGLE)

    def oval(self, s, l, tp, w, h, fill=None, line=None, alpha=None):
        return self.rect(s, l, tp, w, h, fill, line, 0.75, 0.0, alpha, MSO_SHAPE.OVAL)

    def line(self, s, x1, y1, x2, y2, color, w=1.0, dash=None):
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = C(color); cn.line.width = Pt(w)
        if dash:
            try:
                cn.line.dash_style = dash
            except Exception:
                pass
        cn.shadow.inherit = False
        return cn

    def text(self, s, txt, l, tp, w, h, size=18, color="1A1A1A", bold=False, font=None, ea=None,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, italic=False, space_after=4, wrap=True):
        f0 = font if font is not None else self.t["title_font"][0]
        e0 = ea if ea is not None else self.t["title_font"][1]
        tb = s.shapes.add_textbox(Inches(l), Inches(tp), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        for i, ln in enumerate(txt.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(space_after)
            r = p.add_run(); r.text = ln
            f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = C(color); f.name = f0
            _set_ea(r, e0)
        return tb

    def bullets(self, s, items, l, tp, w, h, size=16, color=None, accent=None, gap=10, dot=True):
        color = color or self.t["ink2"]; accent = accent or self.t["accent"]
        tb = s.shapes.add_textbox(Inches(l), Inches(tp), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap); p.line_spacing = 1.15
            rb = p.add_run(); rb.text = ("●  " if dot else "—  ")
            rb.font.size = Pt(size * 0.7); rb.font.color.rgb = C(accent); rb.font.name = self.t["body_font"][0]
            _set_ea(rb, self.t["body_font"][1])
            r = p.add_run(); r.text = it
            r.font.size = Pt(size); r.font.color.rgb = C(color); r.font.name = self.t["body_font"][0]
            _set_ea(r, self.t["body_font"][1])
        return tb

    def pic_ph(self, s, l, tp, w, h, label="图片占位", tone=None):
        tone = tone or self.t["line"]
        shp = self.rect(s, l, tp, w, h, self.t["surface"], tone, 1.0)
        try:
            shp.line.dash_style = 4
        except Exception:
            pass
        self.text(s, label, l, tp + h / 2 - 0.18, w, 0.36, size=12, color=self.t["ink2"],
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                  font=self.t["body_font"][0], ea=self.t["body_font"][1])
        return shp

    def header(self, s, kicker, title, accent=None):
        accent = accent or self.t["accent"]
        self.text(s, kicker.upper(), 0.9, 0.62, 8.5, 0.3, size=11.5, color=accent, bold=True,
                  font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.text(s, title, 0.9, 0.95, 11.6, 0.6, size=28, color=self.t["ink"], bold=True)
        self.rect(s, 0.9, 1.72, 0.62, 0.055, accent)

    def footer(self, s, idx=None):
        self.rect(s, 0.9, SH - 0.62, SW - 1.8, 0.008, self.t["line"])
        self.text(s, self.t["label"] + "  ·  " + self.t["en"], 0.9, SH - 0.55, 7.2, 0.3, size=9, color=self.t["ink2"],
                  font=self.t["body_font"][0], ea=self.t["body_font"][1])
        if idx is not None:
            self.text(s, "%02d" % idx, SW - 1.9, SH - 0.55, 1.0, 0.3, size=9, color=self.t["ink2"],
                      align=PP_ALIGN.RIGHT, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self._page = idx or self._page

    # ---------- 装饰 ----------
    def decor(self, s, variant="cover"):
        d = self.t.get("decor", "min"); a, a2 = self.t["accent"], self.t["accent2"]
        if d == "grid":
            for x in range(1, 13):
                self.rect(s, x, 0, 0.006, SH, self.t["line"])
            self.rect(s, 0.9, 0, 0.05, SH, a)
        elif d == "band":
            self.rect(s, 0, 0, SW, 0.16, a); self.rect(s, 0, 0.16, SW, 0.06, a2)
        elif d == "glow":
            self.oval(s, 9.6, -1.6, 6.2, 6.2, a, alpha=0.16); self.oval(s, -1.8, 4.6, 5.4, 5.4, a2, alpha=0.13)
        elif d == "pixel":
            for i in range(6):
                self.rect(s, 11.4 + (i % 3) * 0.42, 0.7 + (i // 3) * 0.42, 0.34, 0.34,
                          a if i % 2 == 0 else a2, alpha=0.85 if i % 2 == 0 else 0.55)
        elif d == "pop":
            self.oval(s, 10.4, -1.0, 3.2, 3.2, a, alpha=0.9)
            self.oval(s, 11.9, 1.5, 2.1, 2.1, self.t.get("decor3", a2), alpha=0.9)
            self.rect(s, 10.3, 3.2, 1.5, 1.5, a2, alpha=0.85)
        elif d == "glass":
            self.oval(s, 9.0, -1.4, 6.0, 6.0, a, alpha=0.20); self.oval(s, 8.2, 2.6, 7.2, 7.2, a2, alpha=0.14)
        elif d == "term":
            self.text(s, "$", 12.2, 0.6, 0.6, 0.5, size=18, color=a, font=self.t["body_font"][0])
            self.rect(s, 12.75, 0.72, 0.14, 0.28, a)
        elif d == "organic":
            self.oval(s, 10.6, -1.2, 3.6, 3.6, a, alpha=0.20); self.oval(s, 11.6, 2.2, 2.6, 2.6, a2, alpha=0.22)
        elif d == "soft":
            self.oval(s, 9.4, -1.6, 6.0, 6.0, a, alpha=0.14)
        elif d == "rule":
            self.rect(s, 0.9, 0.52, 2.4, 0.02, a)

    # ================= 版式原型 =================
    def cover(self, title="在此输入演示标题", subtitle="在此输入副标题 · 团队 / 姓名",
              meta="YYYY-MM-DD  ·  部门 / 场景", variant="left"):
        s = self.slide(); self.bg(s); self.decor(s, "cover"); a = self.t["accent"]
        if variant == "center":
            self.text(s, title, 1.2, 2.6, 10.9, 1.6, size=46, bold=True, color=self.t["ink"], align=PP_ALIGN.CENTER, spacing=1.02)
            self.rect(s, SW/2 - 0.45, 4.25, 0.9, 0.1, a)
            self.text(s, subtitle, 1.2, 4.5, 10.9, 0.6, size=18, color=self.t["ink2"], align=PP_ALIGN.CENTER, font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.text(s, meta, 1.2, 6.2, 10.9, 0.4, size=12, color=self.t["ink2"], align=PP_ALIGN.CENTER, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        elif variant == "split":
            self.rect(s, 0, 0, 5.4, SH, a, alpha=0.95)
            self.text(s, "SECTION\n01", 0.8, 2.4, 4.0, 1.6, size=40, bold=True, color="FFFFFF", spacing=1.05)
            self.text(s, title, 6.0, 2.5, 6.6, 2.0, size=36, bold=True, color=self.t["ink"], spacing=1.03)
            self.text(s, subtitle, 6.0, 4.6, 6.4, 0.7, size=17, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        else:
            self.rect(s, 0.9, 2.05, 0.9, 0.1, a)
            self.text(s, title, 0.9, 2.4, 10.6, 1.6, size=44, bold=True, color=self.t["ink"], spacing=1.02)
            self.text(s, subtitle, 0.9, 4.15, 10.0, 0.6, size=18, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.text(s, meta, 0.9, 5.9, 8.0, 0.4, size=12, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        return s

    def agenda(self, items=None, title="目录  CONTENTS", variant="num"):
        items = items or ["章节标题一", "章节标题二", "章节标题三", "章节标题四"]
        s = self.slide(); self.bg(s); self.header(s, "Overview", title)
        if variant == "card":
            cw = (SW - 1.8 - 3 * 0.3) / 4
            for i, it in enumerate(items[:4]):
                cx = 0.9 + i * (cw + 0.3)
                self.rrect(s, cx, 2.6, cw, 2.4, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.06))
                self.text(s, "%02d" % (i + 1), cx + 0.35, 2.85, cw - 0.7, 0.5, size=26, bold=True, color=self.t["accent"])
                self.text(s, it, cx + 0.35, 3.7, cw - 0.7, 1.0, size=15, color=self.t["ink"], spacing=1.1)
        else:
            for i, it in enumerate(items):
                cx = 0.9 + (i % 2) * 6.0; cy = 2.35 + (i // 2) * 1.0
                self.text(s, "%02d" % (i + 1), cx, cy, 0.8, 0.5, size=18, bold=True, color=self.t["accent"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
                self.text(s, it, cx + 1.05, cy - 0.02, 4.6, 0.5, size=18, color=self.t["ink"])
                self.line(s, cx + 1.05, cy + 0.56, cx + 5.6, cy + 0.56, self.t["line"], 0.75)
        self.footer(s)
        return s

    def section(self, num="01", title="章节标题", desc="章节说明 · 一句话点题"):
        s = self.slide(); self.bg(s); self.decor(s, "cover")
        self.text(s, num, 0.9, 2.1, 3.0, 1.6, size=96, bold=True, color=self.t["accent"], spacing=0.95)
        self.rect(s, 0.9, 3.85, 0.62, 0.06, self.t["ink"])
        self.text(s, title, 0.9, 4.05, 9.5, 1.0, size=34, bold=True, color=self.t["ink"])
        self.text(s, desc, 0.9, 5.15, 8.6, 0.6, size=15, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def bullets_page(self, kicker, title, items, variant="dot"):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        if variant == "card":
            for i, it in enumerate(items):
                cy = 2.3 + i * 0.86
                self.rrect(s, 0.9, cy, 11.4, 0.7, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.05))
                self.text(s, "%02d" % (i + 1), 1.15, cy + 0.16, 0.7, 0.4, size=13, bold=True, color=self.t["accent"])
                self.text(s, it, 2.0, cy + 0.14, 10.0, 0.45, size=15, color=self.t["ink"])
        elif variant == "num":
            for i, it in enumerate(items[:5]):
                cy = 2.25 + i * 0.86
                self.text(s, "%02d" % (i + 1), 0.9, cy, 1.0, 0.6, size=22, bold=True, color=self.t["accent"])
                self.text(s, it, 2.1, cy + 0.05, 10.0, 0.5, size=16, color=self.t["ink"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
                self.line(s, 2.1, cy + 0.66, 12.4, cy + 0.66, self.t["line"], 0.6)
        else:
            self.bullets(s, items, 0.9, 2.35, 11.4, 4.2, size=17, gap=14)
        self.footer(s)
        return s

    def two_col(self, kicker, title, lt="左侧标题", rt="右侧标题", left=None, right=None, variant="compare"):
        left = left or ["在此输入要点一", "在此输入要点二", "在此输入要点三"]
        right = right or ["在此输入要点一", "在此输入要点二", "在此输入要点三"]
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        for cx, ct, items, ac, tag in [(0.9, lt, left, self.t["accent"], "A"), (6.95, rt, right, self.t["accent2"], "B")]:
            self.rrect(s, cx, 2.3, 5.5, 4.2, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.05))
            self.rect(s, cx, 2.3, 5.5, 0.12, ac)
            self.text(s, tag + " · " + ct, cx + 0.4, 2.6, 4.7, 0.5, size=18, bold=True, color=self.t["ink"])
            self.bullets(s, items, cx + 0.4, 3.35, 4.7, 2.9, size=14.5, accent=ac, gap=11)
        if variant == "vs":
            self.oval(s, SW/2 - 0.42, 4.05, 0.84, 0.84, self.t["ink"])
            self.text(s, "VS", SW/2 - 0.42, 4.05, 0.84, 0.84, size=14, bold=True, color=self.t["bg"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def cards(self, kicker, title, cards_data, cols=4):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        n = len(cards_data); cols = min(cols, max(1, n)); rows = (n + cols - 1) // cols
        gap = 0.3; cw = (SW - 1.8 - (cols - 1) * gap) / cols; ch = (4.25 - (rows - 1) * 0.3) / rows
        for i, (val, lab, desc) in enumerate(cards_data):
            r, c = divmod(i, cols); cx = 0.9 + c * (cw + gap); cy = 2.3 + r * (ch + 0.3)
            self.rrect(s, cx, cy, cw, ch, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.06))
            self.rect(s, cx + 0.35, cy + 0.4, 0.5, 0.06, self.t["accent"])
            self.text(s, val, cx + 0.35, cy + 0.6, cw - 0.7, ch * 0.4, size=28, bold=True, color=self.t["accent"])
            self.text(s, lab, cx + 0.35, cy + ch * 0.64, cw - 0.7, 0.4, size=14, bold=True, color=self.t["ink"])
            self.text(s, desc, cx + 0.35, cy + ch * 0.8, cw - 0.7, ch * 0.3, size=11, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def timeline(self, kicker, title, steps):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        steps = steps or [("阶段一", "说明"), ("阶段二", "说明"), ("阶段三", "说明"), ("阶段四", "说明")]
        n = len(steps); y = 3.35; x0, x1 = 1.35, SW - 1.35
        self.line(s, x0, y, x1, y, self.t["line"], 2.0)
        for i, (st, de) in enumerate(steps):
            cx = x0 + i * (x1 - x0) / max(1, n - 1)
            self.oval(s, cx - 0.16, y - 0.16, 0.32, 0.32, self.t["accent"])
            up = i % 2 == 0
            ty = y - 1.35 if up else y + 0.45
            self.text(s, "%02d" % (i + 1), cx - 0.4, y - 0.62 if up else y + 0.1, 0.8, 0.3, size=11, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)
            self.text(s, st, cx - 1.0, ty, 2.0, 0.4, size=15, bold=True, color=self.t["ink"], align=PP_ALIGN.CENTER)
            self.text(s, de, cx - 1.0, ty + 0.42, 2.0, 0.8, size=11, color=self.t["ink2"], align=PP_ALIGN.CENTER, font=self.t["body_font"][0], ea=self.t["body_font"][1], spacing=1.05)
        self.footer(s)
        return s

    def process(self, kicker, title, nodes, style="chevron"):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        nodes = nodes or ["步骤一", "步骤二", "步骤三", "步骤四"]
        n = len(nodes); gap = 0.25; cw = (SW - 1.8 - (n - 1) * gap) / n
        for i, nd in enumerate(nodes):
            cx = 0.9 + i * (cw + gap)
            fill = self.t["accent"] if i % 2 == 0 else self.t["surface"]
            ink = "FFFFFF" if i % 2 == 0 else self.t["ink"]
            shp = MSO_SHAPE.CHEVRON if style == "chevron" else MSO_SHAPE.ROUNDED_RECTANGLE
            rad = 0.12 if style != "chevron" else 0.0
            self.rect(s, cx, 3.0, cw, 1.5, fill, self.t["line"], 0.75, rad, None, shp)
            self.text(s, nd, cx + 0.22, 3.55, cw - 0.44, 0.5, size=14, bold=True, color=ink, align=PP_ALIGN.CENTER, wrap=True)
            if i < n - 1:
                self.text(s, "›", cx + cw + 0.02, 3.55, gap - 0.04, 0.5, size=18, color=self.t["ink2"], align=PP_ALIGN.CENTER)
        self.footer(s)
        return s

    def table_ph(self, kicker, title, headers, rows):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        cols = len(headers); x0, y0 = 0.9, 2.35; tw = SW - 1.8
        rowh = min(0.7, 4.0 / (len(rows) + 1)); cw = tw / cols
        self.rect(s, x0, y0, tw, rowh, self.t["accent"])
        for j, hd in enumerate(headers):
            self.text(s, hd, x0 + j * cw + 0.2, y0 + rowh * 0.2, cw - 0.4, rowh * 0.6, size=13, bold=True, color="FFFFFF", font=self.t["body_font"][0], ea=self.t["body_font"][1])
        for i, row in enumerate(rows):
            cy = y0 + (i + 1) * rowh
            if i % 2 == 0:
                self.rect(s, x0, cy, tw, rowh, self.t["surface"])
            self.line(s, x0, cy, x0 + tw, cy, self.t["line"], 0.6)
            for j, cell in enumerate(row):
                self.text(s, cell, x0 + j * cw + 0.2, cy + rowh * 0.2, cw - 0.4, rowh * 0.6, size=12.5, color=self.t["ink"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def chart_ph(self, kicker, title, variant="bar"):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        self.rrect(s, 0.9, 2.3, 8.4, 4.25, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.05))
        if variant == "bar":
            for i, v in enumerate([0.55, 0.8, 0.45, 0.95, 0.7, 0.6]):
                bh = 3.2 * v
                self.rect(s, 1.5 + i * 1.25, 6.0 - bh, 0.75, bh, self.t["accent"] if i % 2 == 0 else self.t["accent2"], alpha=0.9)
        elif variant == "line":
            pts = [(1.6, 5.2), (3.0, 4.2), (4.4, 4.7), (5.8, 3.4), (7.2, 3.9), (8.6, 3.0)]
            for i in range(len(pts) - 1):
                self.line(s, pts[i][0], pts[i][1], pts[i + 1][0], pts[i + 1][1], self.t["accent"], 2.5)
            for p in pts:
                self.oval(s, p[0] - 0.07, p[1] - 0.07, 0.14, 0.14, self.t["accent"])
        else:
            self.oval(s, 3.6, 3.0, 2.8, 2.8, self.t["accent"], alpha=0.25)
            self.oval(s, 4.5, 3.9, 1.0, 1.0, self.t["surface"])
            self.text(s, "图", 4.5, 3.9, 1.0, 1.0, size=16, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.text(s, "图表占位 · 在 PowerPoint 中插入图表替换", 1.4, 6.15, 6.4, 0.3, size=10.5, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.text(s, "关键结论 / 解读", 9.7, 2.45, 2.9, 0.4, size=15, bold=True, color=self.t["ink"])
        self.bullets(s, ["在此输入要点一", "在此输入要点二", "在此输入要点三"], 9.7, 3.0, 2.9, 3.2, size=12.5, gap=10)
        self.footer(s)
        return s

    def gallery(self, kicker, title, cols=3, rows=2):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        gap = 0.3; cw = (SW - 1.8 - (cols - 1) * gap) / cols; ch = (4.25 - (rows - 1) * gap) / rows
        for i in range(cols * rows):
            r, c = divmod(i, cols); cx = 0.9 + c * (cw + gap); cy = 2.3 + r * (ch + gap)
            self.pic_ph(s, cx, cy, cw, ch, "作品 %d" % (i + 1))
        self.footer(s)
        return s

    def quote(self, quote_text, author="— 署名 / 职位"):
        s = self.slide(); self.bg(s); self.decor(s, "cover")
        self.text(s, "“", 1.0, 1.6, 2.0, 1.4, size=110, bold=True, color=self.t["accent"])
        self.text(s, quote_text, 1.4, 2.9, 10.5, 2.2, size=26, color=self.t["ink"], spacing=1.2)
        self.rect(s, 1.4, 5.2, 0.62, 0.06, self.t["accent"])
        self.text(s, author, 1.4, 5.4, 8.0, 0.5, size=15, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def team(self, kicker, title, members):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        members = members or [("姓名", "角色")] * 4
        n = len(members); gap = 0.35; cw = (SW - 1.8 - (n - 1) * gap) / n
        for i, (nm, ro) in enumerate(members):
            cx = 0.9 + i * (cw + gap)
            self.rrect(s, cx, 2.5, cw, 3.6, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.06))
            self.oval(s, cx + cw / 2 - 0.65, 2.95, 1.3, 1.3, self.t["line"])
            self.text(s, "头像", cx + cw / 2 - 0.65, 2.95, 1.3, 1.3, size=11, color=self.t["ink2"], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.text(s, nm, cx + 0.2, 4.5, cw - 0.4, 0.4, size=15, bold=True, color=self.t["ink"], align=PP_ALIGN.CENTER)
            self.text(s, ro, cx + 0.2, 4.95, cw - 0.4, 0.4, size=12, color=self.t["accent"], align=PP_ALIGN.CENTER, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def closing(self, title="谢谢观看", sub="在此输入联系方式 / 二维码 / 致谢"):
        s = self.slide(); self.bg(s); self.decor(s, "cover")
        self.text(s, title, 1.2, 2.9, 10.9, 1.2, size=48, bold=True, color=self.t["ink"], align=PP_ALIGN.CENTER)
        self.rect(s, SW/2 - 0.45, 4.3, 0.9, 0.1, self.t["accent"])
        self.text(s, sub, 1.2, 4.6, 10.9, 0.7, size=16, color=self.t["ink2"], align=PP_ALIGN.CENTER, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    # ---------- 新增原型 ----------
    def plans(self, kicker, title, cols_data):
        """方案 / 套餐对比列。"""
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        n = len(cols_data); gap = 0.3; cw = (SW - 1.8 - (n - 1) * gap) / n
        for i, (name, price, items) in enumerate(cols_data):
            cx = 0.9 + i * (cw + gap)
            hi = i == n // 2
            self.rrect(s, cx, 2.35, cw, 4.15, self.t["surface"] if not hi else self.t["accent"], self.t["line"], radius=max(self.t["radius"], 0.06))
            ink = self.t["ink"] if not hi else "FFFFFF"
            sub = self.t["ink2"] if not hi else "FFFFFF"
            self.text(s, name, cx + 0.35, 2.65, cw - 0.7, 0.4, size=16, bold=True, color=ink)
            self.text(s, price, cx + 0.35, 3.1, cw - 0.7, 0.7, size=30, bold=True, color=ink if not hi else "FFFFFF")
            self.bullets(s, items, cx + 0.35, 4.0, cw - 0.7, 2.2, size=12.5, color=sub, accent=self.t["accent"] if not hi else "FFFFFF", gap=9)
        self.footer(s)
        return s

    def faq(self, kicker, title, pairs):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        for i, (q, a) in enumerate(pairs[:4]):
            cy = 2.3 + i * 1.05
            self.rrect(s, 0.9, cy, 11.4, 0.92, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.05))
            self.oval(s, 1.16, cy + 0.24, 0.44, 0.44, self.t["accent"])
            self.text(s, "Q", 1.16, cy + 0.24, 0.44, 0.44, size=13, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.text(s, q, 1.85, cy + 0.16, 10.1, 0.4, size=14, bold=True, color=self.t["ink"])
            self.text(s, a, 1.85, cy + 0.54, 10.1, 0.34, size=11.5, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def glossary(self, kicker, title, terms):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        for i, (term, dfn) in enumerate(terms[:6]):
            cy = 2.3 + i * 0.72
            self.text(s, term, 0.9, cy, 3.2, 0.5, size=15, bold=True, color=self.t["accent"])
            self.text(s, dfn, 4.3, cy, 8.2, 0.5, size=13, color=self.t["ink"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.line(s, 0.9, cy + 0.58, 12.4, cy + 0.58, self.t["line"], 0.6)
        self.footer(s)
        return s

    def compare_table(self, kicker, title, headers, rows, marks):
        """带标记的对比矩阵。"""
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        cols = len(headers); x0, y0 = 0.9, 2.35; tw = SW - 1.8
        rowh = min(0.72, 4.2 / (len(rows) + 1)); cw = tw / cols
        self.rect(s, x0, y0, tw, rowh, self.t["ink"])
        for j, hd in enumerate(headers):
            self.text(s, hd, x0 + j * cw + 0.2, y0 + rowh * 0.22, cw - 0.4, rowh * 0.6, size=13, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        for i, row in enumerate(rows):
            cy = y0 + (i + 1) * rowh
            if i % 2 == 0:
                self.rect(s, x0, cy, tw, rowh, self.t["surface"])
            self.line(s, x0, cy, x0 + tw, cy, self.t["line"], 0.6)
            for j, cell in enumerate(row):
                if j == 0:
                    self.text(s, cell, x0 + 0.2, cy + rowh * 0.22, cw - 0.4, rowh * 0.6, size=12.5, color=self.t["ink"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
                else:
                    mk = marks[i][j - 1] if i < len(marks) and (j - 1) < len(marks[i]) else 1
                    col = self.t["accent"] if mk else self.t["line"]
                    self.oval(s, x0 + j * cw + cw / 2 - 0.11, cy + rowh * 0.5 - 0.11, 0.22, 0.22, col)
        self.footer(s)
        return s

    def roadmap(self, kicker, title, lanes):
        """多泳道路线图。"""
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        n = len(lanes); lh = 3.6 / n
        for li, (lname, blocks) in enumerate(lanes):
            y = 2.35 + li * lh
            self.text(s, lname, 0.9, y + 0.24, 1.5, 0.4, size=13, bold=True, color=self.t["ink"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.line(s, 2.5, y + lh / 2, 12.4, y + lh / 2, self.t["line"], 1.0)
            for bi, (bw, lab) in enumerate(blocks):
                bx = 2.6 + bi * 2.4
                self.rrect(s, bx, y + 0.12, bw * 1.9, lh - 0.34, self.t["accent"] if li % 2 == 0 else self.t["accent2"], None, radius=0.3, alpha=0.9)
                self.text(s, lab, bx + 0.12, y + 0.16, bw * 1.9 - 0.24, lh - 0.42, size=10.5, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, wrap=True)
        self.footer(s)
        return s

    def stat_chart(self, kicker, title, big, cap):
        """大指标 + 迷你图表组合。"""
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        self.text(s, big, 0.9, 2.4, 5.0, 1.6, size=72, bold=True, color=self.t["accent"])
        self.text(s, cap, 0.95, 4.1, 5.0, 0.8, size=15, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1], spacing=1.15)
        self.rrect(s, 6.4, 2.4, 5.9, 3.9, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.05))
        for i, v in enumerate([0.4, 0.65, 0.5, 0.85, 0.6, 0.95, 0.7]):
            bh = 2.8 * v
            self.rect(s, 6.9 + i * 0.76, 5.9 - bh, 0.48, bh, self.t["accent"] if i % 2 == 0 else self.t["accent2"], alpha=0.9)
        self.footer(s)
        return s

    def bignumber(self, kicker, number, caption, note="在此输入补充说明"):
        s = self.slide(); self.bg(s); self.decor(s, "cover")
        self.text(s, kicker.upper(), 0.9, 0.9, 8.0, 0.35, size=12, color=self.t["accent"], bold=True, font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.text(s, number, 0.9, 1.9, 11.0, 2.2, size=120, bold=True, color=self.t["ink"])
        self.text(s, caption, 1.0, 4.35, 10.0, 0.7, size=22, color=self.t["accent"], bold=True)
        self.text(s, note, 1.0, 5.2, 9.0, 0.6, size=14, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def checklist(self, kicker, title, items):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        for i, it in enumerate(items[:6]):
            cy = 2.3 + i * 0.72
            self.rect(s, 0.95, cy + 0.05, 0.34, 0.34, None, self.t["accent"], 1.4)
            self.text(s, "✓", 0.95, cy + 0.03, 0.34, 0.38, size=13, bold=True, color=self.t["accent"], align=PP_ALIGN.CENTER)
            self.text(s, it, 1.7, cy, 10.6, 0.5, size=14.5, color=self.t["ink"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.line(s, 1.7, cy + 0.56, 12.4, cy + 0.56, self.t["line"], 0.6)
        self.footer(s)
        return s

    def quote_wall(self, kicker, title, quotes):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        n = len(quotes); cols = 3; rows = (n + cols - 1) // cols
        gap = 0.28; cw = (SW - 1.8 - (cols - 1) * gap) / cols; ch = (4.25 - (rows - 1) * gap) / rows
        for i, q in enumerate(quotes[:6]):
            r, c = divmod(i, cols); cx = 0.9 + c * (cw + gap); cy = 2.3 + r * (ch + gap)
            self.rrect(s, cx, cy, cw, ch, self.t["surface"], self.t["line"], radius=max(self.t["radius"], 0.06))
            self.text(s, "“", cx + 0.28, cy + 0.12, 1.0, 0.6, size=34, bold=True, color=self.t["accent"])
            self.text(s, q, cx + 0.32, cy + 0.7, cw - 0.64, ch - 1.0, size=13, color=self.t["ink"], spacing=1.2)
        self.footer(s)
        return s

    def steps_vertical(self, kicker, title, steps):
        s = self.slide(); self.bg(s); self.header(s, kicker, title)
        n = min(5, len(steps)); x = 1.5
        self.line(s, x, 2.5, x, 2.5 + (n - 1) * 0.85, self.t["line"], 1.6)
        for i, (st, de) in enumerate(steps[:n]):
            cy = 2.4 + i * 0.85
            self.oval(s, x - 0.19, cy + 0.02, 0.38, 0.38, self.t["accent"])
            self.text(s, str(i + 1), x - 0.19, cy + 0.02, 0.38, 0.38, size=12, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font=self.t["body_font"][0], ea=self.t["body_font"][1])
            self.text(s, st, x + 0.5, cy - 0.05, 4.0, 0.4, size=15, bold=True, color=self.t["ink"])
            self.text(s, de, x + 0.5, cy + 0.32, 9.8, 0.4, size=12.5, color=self.t["ink2"], font=self.t["body_font"][0], ea=self.t["body_font"][1])
        self.footer(s)
        return s

    def save(self, path):
        base = getattr(self, "_seed", 0)
        for i, s in enumerate(self.prs.slides):
            spec = _anim.TRANSITIONS[(base + i) % len(_anim.TRANSITIONS)]
            _anim.apply_to_slide(s, base + i, spec, SW, SH)
        self.prs.save(path)
        return path
